"""测试用材料构造工具（合成数据）。"""

from __future__ import annotations

import io

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.comments import Comment


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None, note_after_table: str | None = None,
              comment: str | None = None, hidden: str | None = None, tracked_delete: str | None = None,
              title_prop: str | None = None) -> bytes:
    d = docx.Document()
    for t in paragraphs:
        d.add_paragraph(t)
    if table:
        tb = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                tb.cell(r, c).text = val
        if note_after_table:
            d.add_paragraph(note_after_table)
    if comment:
        p = d.add_paragraph("需要批注的段落")
        d.add_comment(p.runs, text=comment, author="审核人")
    if hidden:
        p = d.add_paragraph()
        r = p.add_run(hidden)
        rpr = r._element.get_or_add_rPr()
        v = OxmlElement("w:vanish")
        rpr.append(v)
    if tracked_delete:
        p = d.add_paragraph("保留文字")
        dl = OxmlElement("w:del")
        dl.set(qn("w:id"), "1")
        dl.set(qn("w:author"), "x")
        r = OxmlElement("w:r")
        t = OxmlElement("w:delText")
        t.text = tracked_delete
        r.append(t)
        dl.append(r)
        p._p.append(dl)
    if title_prop:
        d.core_properties.title = title_prop
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx(rows: list[list], sheet: str = "经费测算", hidden_sheet: list[list] | None = None,
              comment_at: tuple[str, str] | None = None, formulas: dict[str, str] | None = None,
              hide_row: int | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    for coord, f in (formulas or {}).items():
        ws[coord] = f
    if comment_at:
        ws[comment_at[0]].comment = Comment(comment_at[1], "审核人")
    if hide_row:
        ws.row_dimensions[hide_row].hidden = True
    if hidden_sheet:
        hs = wb.create_sheet("隐藏数据")
        for r in hidden_sheet:
            hs.append(r)
        hs.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_numbered_docx(items: list[tuple[int, str]], plain_before: list[str] | None = None, plain_after: list[str] | None = None) -> bytes:
    """Word 自动编号的文稿：items 为 (层级, 文字)，层级 0～3 分别编为“一、”“（一）”“1.”“（1）”，序号不在段落文字中。"""
    d = docx.Document()
    numbering = d.part.numbering_part.element
    absn = OxmlElement("w:abstractNum")
    absn.set(qn("w:abstractNumId"), "90")
    for ilvl, (fmt, text) in enumerate((("chineseCounting", "%1、"), ("chineseCounting", "（%2）"), ("decimal", "%3."), ("decimal", "（%4）"))):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        for tag, val in (("w:start", "1"), ("w:numFmt", fmt), ("w:lvlText", text)):
            e = OxmlElement(tag)
            e.set(qn("w:val"), val)
            lvl.append(e)
        absn.append(lvl)
    numbering.insert(0, absn)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "90")
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), "90")
    num.append(ref)
    numbering.append(num)
    for t in plain_before or []:
        d.add_paragraph(t)
    for level, text in items:
        p = d.add_paragraph(text)
        ppr = p._p.get_or_add_pPr()
        numpr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl")
        il.set(qn("w:val"), str(level))
        ni = OxmlElement("w:numId")
        ni.set(qn("w:val"), "90")
        numpr.append(il)
        numpr.append(ni)
        ppr.append(numpr)
    for t in plain_after or []:
        d.add_paragraph(t)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
