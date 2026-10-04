"""各种版式的编排与实际渲染核验（GB/T 9704—2012 第 10 章特定格式、题注与长标题回行、简报与事务文书）。

* 信函格式（10.1）：发文机关标志上边缘距上页边 30mm；标志下 4mm 处上粗下细、距下页边 20mm 处上细下粗的
  红色双线，线长 170mm 居中；份号等要素居左、发文字号居右，距第一条双线 7/8 个 3 号字高；首页不显示页码；
  版记不加印发机关、印发日期与分隔线；
* 命令（令）格式（10.2）：发文机关标志上边缘距版心上边缘 20mm；签名章（7.3.5.3）右空四字，成文日期在其下空一行；
* 含数字的长标题按实际渲染宽度回行，题注均衡回行，都不留下单字或短尾；
* 简报报头与不设版头的事务文书（实务）。

渲染测试需要 LibreOffice 与 poppler，缺少时跳过。
"""

import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document

from gongwen.layout.docx_compiler import compile_docx
from gongwen.layout.pipeline import layout_document
from gongwen.layout.profile import LayoutProfile, render_width_chars, split_balanced, split_title
from gongwen.layout.render_check import check_rendering, double_lines, pdf_pages, red_bands
from gongwen.schemas.ir import Block, DocumentIR, Header, Imprint, Sentence, Signature

MM = 25.4 / 72
ISSUER = "示例市卫生健康委员会"
PARA = "各区要高度重视基层医疗示范点建设工作，确保今年底前完成建设任务，按月报送进展情况，及时协调解决存在的问题。"
SPEECH = "在2026年基层医疗示范点建设推进会上的讲话"
RESOLUTION_NOTE = "（2026年3月10日示例市第十七届人民代表大会第五次会议通过）"

needs_render = pytest.mark.skipif(not (shutil.which("soffice") and shutil.which("pdftotext") and shutil.which("pdftoppm")), reason="需要 LibreOffice 与 poppler 做实际渲染")


def _blocks(n: int) -> list[Block]:
    return [Block(bid=f"b{i}", kind="paragraph", sentences=[Sentence(sid=f"s{i}", text=PARA)]) for i in range(n)]


def _letter(n: int = 3, **header) -> DocumentIR:
    return DocumentIR(
        doc_id="L1", matter_id="m", genre="函", format_type="letter", blocks=_blocks(n),
        header=Header(organ_mark=ISSUER, doc_number="示卫函〔2026〕5号", **header),
        title=f"{ISSUER}关于商请支持示范点建设的函", recipients=["示例市财政局"],
        signature=Signature(organs=[ISSUER], date="2026年8月1日"),
        imprint=Imprint(cc=["示例市发展和改革委员会"], printer="【待确认印发机关】", print_date="【待印发时填写】"),
    )


def _command(n: int = 3) -> DocumentIR:
    return DocumentIR(
        doc_id="C1", matter_id="m", genre="命令（令）", format_type="command", blocks=_blocks(n),
        header=Header(organ_mark="示例市人民政府令", doc_number="第【待编号：令号由办理流程确定】号"),
        signature=Signature(organs=["示例市人民政府"], seal_mode="signature_stamp", signer_title="市长", date="2026年8月1日"),
    )


def _brief(**meta) -> DocumentIR:
    return DocumentIR(
        doc_id="B1", matter_id="m", format_type="brief", blocks=_blocks(3), meta=meta,
        header=Header(organ_mark="工作简报", doc_number="第【待编号：期号由办理流程确定】期"),
        title="示例市扎实推进基层医疗示范点建设", signature=Signature(organs=[], seal_mode="none", date=""),
        imprint=Imprint(cc=["示例市财政局"], printer="示例市卫生健康委员会办公室", print_date="2026年8月1日"),
    )


def _speech() -> DocumentIR:
    return DocumentIR(
        doc_id="P1", matter_id="m", format_type="plain", blocks=_blocks(3), header=Header(organ_mark="", doc_number=""),
        title=SPEECH, title_note="副市长　【待补：讲话人姓名】", salutation="同志们：", signature=Signature(organs=[], seal_mode="none", date=""),
    )


def _render(ir: DocumentIR, out: Path, profile: LayoutProfile | None = None, check_profile: LayoutProfile | None = None):
    profile = profile or LayoutProfile.load()
    path, fonts = compile_docx(ir, out / f"{ir.doc_id}.docx", profile)
    info, checks = check_rendering(path, ir, check_profile or profile, fonts, out)
    return Path(info.pdf_path), pdf_pages(Path(info.pdf_path)), {c.rule_id: c for c in checks}, checks


def _line(lines, text: str):
    return next(l for l in lines if text in l.text)


# ------------------------------------------------------------------ 回行（不需要渲染）
def test_rendered_width_counts_digits_wider_than_half_a_character():
    assert render_width_chars("基层医疗") == 4
    assert 4.8 < render_width_chars("在2026年") < 5.1  # 实测 4.95 字：数字约 0.64 字，与汉字相邻处另有自动间距
    assert render_width_chars(SPEECH) > 20  # 按半角 0.5 字计恰为 20 字，实际排不下 156mm


def test_long_title_with_digits_never_leaves_an_orphan():
    lines = split_title(SPEECH, 20)
    assert len(lines) == 2 and "".join(lines) == SPEECH, lines
    assert all(render_width_chars(x) <= 20 and len(x) >= 6 for x in lines), lines
    two = split_title(f"{ISSUER}关于2026年度基层医疗示范点建设情况的报告", 20, ISSUER)
    assert all(len(x) >= 6 for x in two) and not any(x.endswith("2026") for x in two), two  # 不拆“2026年”，不留“的报告”短尾


def test_title_note_is_split_into_balanced_lines():
    lines = split_balanced(RESOLUTION_NOTE, 28)
    assert len(lines) == 2 and "".join(lines) == RESOLUTION_NOTE
    assert min(len(x) for x in lines) >= 10 and all(render_width_chars(x) <= 28 for x in lines), lines
    assert any("第十七届" in x for x in lines)  # 序数词组不拆开
    assert split_balanced("副市长　【待补：讲话人姓名】", 28) == ["副市长　【待补：讲话人姓名】"]


# ------------------------------------------------------------------ DOCX 结构（不需要渲染）
def test_letter_docx_uses_page_anchored_frames_and_blank_first_footer(tmp_path):
    path, _ = compile_docx(_letter(copy_no="000001"), tmp_path / "l.docx")
    z = zipfile.ZipFile(path)
    xml = z.read("word/document.xml").decode("utf-8")
    assert xml.count('w:vAnchor="page"') >= 4  # 标志、两条双线、份号与发文字号
    assert "<w:titlePg/>" in xml
    doc = Document(str(path))
    first_footer = "".join(p._p.xml for p in doc.sections[0].first_page_footer.paragraphs)
    assert "PAGE" not in first_footer  # 首页不显示页码（10.1）
    texts = [p.text for p in doc.paragraphs]
    assert "000001\t示卫函〔2026〕5号" in texts  # 份号居左、发文字号居右，同排一行
    cells = " ".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "抄送" in cells and "印发" not in cells and "待确认印发机关" not in cells  # 版记不加印发机关和印发日期


def test_brief_header_and_plain_title_placement(tmp_path):
    path, _ = compile_docx(_brief(), tmp_path / "b.docx")
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert "工作简报" in texts and "第【待编号：期号由办理流程确定】期" in texts
    assert "【待补：编印单位】\t【待补：编印日期】" in texts
    assert not doc.tables  # 简报不设版记
    path, _ = compile_docx(_brief(brief_issuer="示例市卫生健康委员会办公室", brief_date="2026年3月10日"), tmp_path / "b2.docx")
    assert "示例市卫生健康委员会办公室\t2026年3月10日" in [p.text for p in Document(str(path)).paragraphs]
    path, _ = compile_docx(_speech(), tmp_path / "p.docx")
    first = Document(str(path)).paragraphs[0]
    assert first.text == split_title(SPEECH, 20)[0]  # 事务文书不设版头，标题前没有空行


# ------------------------------------------------------------------ 实际渲染
@needs_render
def test_letter_format_measures_per_gbt9704_10_1(tmp_path):
    ir = _letter(14, copy_no="000001", secrecy="机密★1年", urgency="特急")
    pdf, pages, by, checks = _render(ir, tmp_path)
    first = pages[0][2]
    mark = _line(first, ISSUER)
    assert abs(mark.y0_mm - 30) <= 1, mark.y0_mm  # 修正前约 35.3mm
    for rid in ("LAY-LETTER-MARK", "LAY-LETTER-RULES", "LAY-LETTER-DOCNO", "LAY-PAGENO"):
        assert by[rid].status == "pass", (rid, by[rid].actual, by[rid].note)
    assert not any(c.status in ("fail", "warn") for c in checks), [(c.rule_id, c.actual) for c in checks if c.status != "pass"]
    bands = red_bands(pdf)
    (t1, t2), (b1, b2) = double_lines(bands)
    ink = [b for b in bands if b.height > 2][0]
    assert abs(t1.y0 - ink.y1 - 4) <= 1  # 标志下 4mm
    assert t1.height > 2 * t2.height and b2.height > 2 * b1.height  # 上粗下细、上细下粗
    assert abs(297 - b2.y1 - 20) <= 0.5  # 距下页边 20mm
    for a, b in ((t1, t2), (b1, b2)):
        assert abs(b.x1 - a.x0 - 170) <= 1 and abs((a.x0 + b.x1) / 2 - 106) <= 1  # 170mm，以版心为准居中
    docno, copy_no = _line(first, "示卫函"), _line(first, "000001")
    assert abs(docno.y0_mm - t2.y1 - 16 * 7 / 8 * MM) <= 0.5  # 3 号字高的 7/8
    assert abs(docno.x1 * MM - 184) <= 0.5 and abs(copy_no.x0_mm - 28) <= 0.5 and abs(copy_no.y0_mm - docno.y0_mm) <= 0.5
    urgency = _line(first, "特急")
    assert urgency.y0_mm > _line(first, "机密").y0_mm > docno.y0_mm  # 份号、密级、紧急程度自上而下
    title = [l for l in first if l.text == ISSUER][1]  # 标题首行（回行为“示例市卫生健康委员会／关于……的函”）
    assert 2.5 * 10.2 < title.y0_mm - urgency.y0_mm < 3.5 * 10.2  # 标题与其上最后一个要素相距二行
    assert not any(l.text.startswith("—") for l in first)  # 首页不显示页码
    assert len(pages) >= 2 and any(l.text.startswith("—") for l in pages[1][2]) and not red_bands(pdf, 2)  # 第二面有页码、无双线


@needs_render
def test_letter_checks_report_misplaced_mark_and_first_page_number(tmp_path):
    """检查本身能发现修正前的问题：标志不在 30mm 处、首页显示页码。"""
    wrong = LayoutProfile.load(overrides={"letter": {"organ_mark_top_from_page_mm": 36, "first_page_number": True}})
    _, _, by, _ = _render(_letter(), tmp_path, profile=wrong, check_profile=LayoutProfile.load())
    assert by["LAY-LETTER-MARK"].status == "warn" and by["LAY-PAGENO"].status == "fail"


@needs_render
def test_general_mark_unchanged_and_plain_title_starts_at_type_area_top(tmp_path):
    general = DocumentIR(
        doc_id="G1", matter_id="m", format_type="general", blocks=_blocks(2), header=Header(organ_mark=f"{ISSUER}文件", doc_number="示卫发〔2026〕5号"),
        title="关于加强基层医疗示范点建设的通知", recipients=["各区卫生健康局"], signature=Signature(organs=[ISSUER], date="2026年8月1日"),
        imprint=Imprint(printer=f"{ISSUER}办公室", print_date="2026年8月1日"),
    )
    _, _, by, _ = _render(general, tmp_path / "g")
    assert by["LAY-MARK"].status == "pass" and abs(float(by["LAY-MARK"].actual.split("mm")[0]) - 35) <= 0.5
    assert not any(k.startswith("LAY-LETTER") for k in by)
    _, pages, by, checks = _render(_speech(), tmp_path / "p")
    lines = pages[0][2]
    title = [l for l in lines if l.text and l.text in SPEECH]
    assert "".join(l.text for l in title) == SPEECH and len(title) == 2 and min(len(l.text) for l in title) >= 6  # 修正前“话”单独回行
    assert abs(title[0].y0_mm - 37) <= 3  # 版心第一行起排，前面没有空行（修正前约 58mm）
    assert "LAY-MARK" not in by and not any(c.status in ("fail", "warn") for c in checks)  # 无版头，不按红头文件核验
    note = _line(lines, "副市长")
    assert note.y0_mm > title[-1].y0_mm


@needs_render
def test_title_note_renders_in_balanced_lines(tmp_path):
    ir = _speech()
    ir.title, ir.title_note, ir.salutation = "示例市人民代表大会关于2026年国民经济和社会发展计划的决议", RESOLUTION_NOTE, ""
    _, pages, _, _ = _render(ir, tmp_path)
    lines = [l for l in pages[0][2] if l.text and l.text in RESOLUTION_NOTE]
    assert "".join(l.text for l in lines) == RESOLUTION_NOTE and len(lines) == 2, [l.text for l in lines]
    assert min(len(l.text) for l in lines) >= 10  # 修正前“通过）”单独成行
    for l in lines:  # 居中
        assert abs((l.x0 + l.x1) / 2 * MM - 106) <= 1.5


@needs_render
def test_brief_header_renders_without_signature_or_imprint(tmp_path):
    pdf, pages, by, checks = _render(_brief(), tmp_path)
    first = pages[0][2]
    mark, issue, issuer, date, title = (_line(first, t) for t in ("工作简报", "期号", "编印单位", "编印日期", "示例市扎实推进"))
    assert by["LAY-MARK"].status == "pass" and by["LAY-MARK"].level == "实务"
    assert mark.y0 < issue.y0 < issuer.y0 < title.y0 and abs(issuer.y0 - date.y0) < 1
    assert abs(issuer.x0_mm - 28) <= 0.5 and abs(date.x1 * MM - 184) <= 0.5  # 左编印单位、右日期
    rule = [b for b in red_bands(pdf) if b.height < 2 and b.x1 - b.x0 > 150]
    assert len(rule) == 1 and issuer.y1 * MM < rule[0].y0 < title.y0_mm  # 红色分隔线在编印单位行与标题之间
    text = "".join(l.text for _, _, ls in pages for l in ls)
    assert "抄送" not in text and "印发" not in text  # 不设版记
    assert "LAY-SIGNATURE" not in by and "LAY-IMPRINT" not in by and not any(c.status in ("fail", "warn") and c.rule_id != "FONT" for c in checks)


@needs_render
def test_command_mark_and_signature_stamp(tmp_path):
    _, pages, by, _ = _render(_command(), tmp_path)
    assert by["LAY-MARK"].status == "pass" and "10.2" in by["LAY-MARK"].clause
    assert abs(float(by["LAY-MARK"].actual.split("mm")[0]) - 20) <= 0.5
    lines = pages[0][2]
    title, stamp, date = _line(lines, "市长"), _line(lines, "【签名章】"), _line(lines, "2026年8月1日")
    pitch = LayoutProfile.load().char_pitch_pt * MM
    assert abs(184 - stamp.x1 * MM - 4 * pitch) <= 0.6  # 签名章右空四字（7.3.5.3）
    assert abs(stamp.x0_mm - title.x1 * MM - 2 * pitch) <= 0.6  # 签名章左空二字标注签发人职务
    assert abs(date.x1 * MM - stamp.x1 * MM) <= 0.6 and abs(date.y0_mm - stamp.y0_mm - 2 * 28.95 * MM) <= 0.5  # 下空一行右空四字
    assert by["LAY-SIGNATURE"].status == "pass"


@needs_render
def test_signature_stamp_alone_on_last_page_is_reported(tmp_path):
    """签名章页没有正文时 LAY-SIGNATURE 不通过：以“签发人职务　【签名章】”行（不是发文机关名称）为署名行。"""
    ir = _command()
    profile = LayoutProfile.load()
    path, fonts = compile_docx(ir, tmp_path / "c.docx", profile)
    doc = Document(str(path))
    for p in doc.paragraphs:
        p.paragraph_format.keep_with_next = False
    next(p for p in doc.paragraphs if "签名章" in p.text).paragraph_format.page_break_before = True
    doc.save(str(path))
    _, checks = check_rendering(path, ir, profile, fonts, tmp_path)
    sig = next(c for c in checks if c.rule_id == "LAY-SIGNATURE")
    assert sig.status == "fail" and sig.actual.startswith("第2面"), sig.actual
    # 正常编排：首页排不下时正文末段随签名章转到下一面，签名章页有正文
    rep = layout_document(_command(5), tmp_path / "pipe")
    sig = next(c for c in rep.checks if c.rule_id == "LAY-SIGNATURE")
    assert sig.status == "pass" and sig.actual == "第2面有正文", sig.actual
