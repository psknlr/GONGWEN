"""公文模板（可自行调整）：合并与校验、偏离国标逐项列出、单位信息、命令行、Word 模板导出、按模板排版与核验。

渲染相关的测试需要 LibreOffice 与 poppler，缺少时跳过。
"""

import json
import shutil
import zipfile

import pytest
from docx import Document
from test_engine_e2e import make_engine, run_to_review

from gongwen.cli.main import main as cli
from gongwen.layout.dotx import DOCUMENT_CT, TEMPLATE_CT, export_dotx
from gongwen.layout.pipeline import layout_document
from gongwen.layout.profile import LayoutProfile
from gongwen.layout.templates import Template, TemplateError, TemplateStore, apply_unit, sample_ir
from gongwen.orchestrator import default_user
from gongwen.schemas.common import Clearance

needs_render = pytest.mark.skipif(not (shutil.which("soffice") and shutil.which("pdftotext") and shutil.which("pdftoppm")), reason="需要 LibreOffice 与 poppler 做实际渲染")


def _errors(data: dict) -> list[str]:
    with pytest.raises(TemplateError) as ei:
        Template.from_dict({"name": "t", **data})
    return ei.value.errors


# ------------------------------------------------------------------ 校验
def test_validation_rejects_unknown_keys_and_invalid_values_with_chinese_messages():
    assert "是否要写 elements.title" in _errors({"elements": {"titel": {"size": "小二"}}})[0]
    assert "未知设置项：colour" in _errors({"colour": "red"})[0]
    assert "应为字号名之一" in _errors({"elements": {"title": {"size": "小五"}}})[0]
    assert "6 位十六进制颜色" in _errors({"elements": {"organ_mark": {"color": "red"}}})[0]
    assert "应在 5～100mm 之间" in _errors({"margins": {"top_mm": 300}})[0]
    assert "应为字体类别之一" in _errors({"elements": {"title": {"font": "Arial"}}})[0]
    assert "标准注记" in _errors({"elements": {"title": {"strength": "推荐"}}})[0]
    assert "基础配置档" in _errors({"base": "其他"})[0]
    # 版心与页边距不一致：给出应有的值
    err = _errors({"margins": {"top_mm": 37, "bottom_mm": 20}})[0]
    assert "≠ 纸张高 297mm" in err and "应为 35mm" in err
    # 行数 × 行距超出版心
    assert "超出版心高" in _errors({"grid": {"lines_per_page": 22, "line_pt": 30}})[0]
    with pytest.raises(TemplateError, match="模板名"):
        Template.from_dict({"name": "../x"})
    # 中文类别名可用作字体类别
    tpl = Template.from_dict({"name": "t", "elements": {"levels": {2: "黑体"}}})
    assert tpl.overrides["elements"]["levels"][2] == "heiti"


def test_deviations_are_listed_with_clause_and_strength():
    tpl = Template.from_dict({
        "name": "本单位",
        "fonts": {"fangsong": {"name": "仿宋"}},
        "margins": {"top_mm": 37.5},
        "elements": {"title": {"size": "小二"}, "organ_mark": {"color": "#C00000", "top_from_type_area_mm": 30}, "red_rule": {"width_pt": 1.5}},
    })
    by = {c.path: c for c in tpl.changes()}
    dev = {c.path for c in tpl.deviations()}
    assert dev == {"elements.title.size", "elements.organ_mark.color", "elements.organ_mark.top_from_type_area_mm"}
    assert (by["elements.title.size"].clause, by["elements.title.size"].strength) == ("7.3.1", "一般")
    assert (by["elements.organ_mark.top_from_type_area_mm"].clause, by["elements.organ_mark.top_from_type_area_mm"].strength) == ("7.2.4", "规定")
    assert by["elements.organ_mark.color"].strength == "推荐"
    # 字库名、线宽属实务；天头 37.5mm 在 ±1mm 允许误差内：都不算偏离国标，但列出
    assert by["fonts.fangsong.name"].kind.startswith("实务") and by["elements.red_rule.width_pt"].kind.startswith("实务")
    assert by["margins.top_mm"].kind == "国标允许误差内"
    # 随之推算：下白边、标题每行字数
    assert by["margins.bottom_mm"].derived and by["margins.bottom_mm"].value == 34.5
    assert by["elements.title.max_chars_per_line"].derived and by["elements.title.max_chars_per_line"].value == 24
    assert "标题字号：二号 → 小二（GB/T 9704—2012 7.3.1，一般）" in by["elements.title.size"].describe()


def test_effective_values_are_derived_to_fit_type_area():
    data = Template.from_dict({"name": "t", "grid": {"lines_per_page": 23}, "elements": {"body": {"size": "小三"}}}).effective_data()
    assert data["grid"]["line_pt"] * 23 <= 225 / 25.4 * 72 and data["grid"]["line_pt"] == 27.7  # 向下取整到缇
    pitch = 15 + data["grid"]["char_spacing_twips"] / 20
    assert 28 * pitch <= 156 / 25.4 * 72 + 0.01  # 每行仍排 28 字
    base = Template.from_dict({"name": "t"}).effective_data()
    assert base == LayoutProfile.load().data  # 空模板与基础配置档完全相同


def test_store_precedence_builtin_protection_and_update(tmp_path):
    store = TemplateStore(tmp_path / "data", tmp_path)
    assert {"gbt9704-2012", "windows-fonts"} <= set(store.names())
    with pytest.raises(ValueError, match="内置模板"):
        store.update("windows-fonts", ["elements.title.size=小二"])
    with pytest.raises(ValueError, match="内置模板"):
        store.save(Template(name="windows-fonts"))
    with pytest.raises(ValueError):
        store.delete("gbt9704-2012")
    tpl = store.new("本单位", from_="windows-fonts")
    assert tpl.overrides["fonts"]["fangsong"]["name"] == "仿宋"
    tpl = store.update("本单位", ["elements.title.size=小二", "elements.organ_mark.color=#C00000", "unit.cc=[示例市财政局, 示例市人民政府办公室]"])
    assert tpl.overrides["elements"]["organ_mark"]["color"] == "C00000" and tpl.unit["cc"] == ["示例市财政局", "示例市人民政府办公室"]
    tpl = store.update("本单位", unsets=["elements.title.size"])
    assert "title" not in tpl.overrides.get("elements", {})
    with pytest.raises(TemplateError):
        store.update("本单位", ["elements.title.size=小五"])
    assert store.get("本单位").overrides["elements"]["organ_mark"]["color"] == "C00000"  # 不合规的修改未写入
    # 工作区 templates/ 中的模板可读不可改；数据目录中的同名模板优先
    (tmp_path / "templates").mkdir()
    (tmp_path / "templates" / "共享.yaml").write_text("name: 共享\nelements: {title: {size: 小二}}\n", encoding="utf-8")
    assert store.get("共享").kind == "工作区"
    with pytest.raises(ValueError, match="工作区"):
        store.update("共享", ["description=x"])
    store.delete("本单位")
    assert "本单位" not in store.names()


def test_unit_info_fills_only_missing_elements():
    ir = sample_ir()
    ir.header.organ_mark, ir.header.doc_number, ir.imprint.printer, ir.imprint.cc = "", "【待编号：发文字号由办理流程确定】", "", []
    unit = {"organ_mark": "示例市卫生健康委员会文件", "doc_number_prefix": "示卫发", "printer": "示例市卫生健康委员会办公室", "cc": ["示例市财政局"]}
    new, filled = apply_unit(ir, unit)
    assert new.header.organ_mark == "示例市卫生健康委员会文件" and ir.header.organ_mark == ""  # 不改动原文稿
    assert new.header.doc_number == "示卫发〔【待补：年份】〕【待编号】号"  # 年份与序号不猜测
    assert new.imprint.printer == "示例市卫生健康委员会办公室" and new.imprint.cc == ["示例市财政局"]
    assert set(filled) == {"发文机关标志", "发文字号代字", "印发机关", "抄送机关（模板默认）"}
    kept, filled = apply_unit(sample_ir(), unit)  # 已有内容不覆盖
    assert kept.header.doc_number == "示卫发〔2026〕1号" and kept.imprint.cc == ["示例市人民政府办公室"] and "发文字号代字" not in filled


def test_page_number_style_and_letter_brief_unit_fields(tmp_path):
    from gongwen.layout.docx_compiler import compile_docx

    tpl = Template.from_dict({"name": "t", "elements": {"page_number": {"style": "plain"}}, "unit": {"letter_organ_mark": "示例市卫生健康委员会", "brief_name": "卫生健康工作简报", "brief_issuer": "示例市卫生健康委员会办公室"}})
    assert [c.strength for c in tpl.deviations()] == ["一般"]  # 页码不加一字线：偏离 7.5（一般）
    path, _ = compile_docx(sample_ir(), tmp_path / "a.docx", tpl.profile())
    footer = "".join(p.text for p in Document(str(path)).sections[0].footer.paragraphs)
    assert "—" not in footer
    letter = sample_ir("letter")
    letter.header.organ_mark = ""
    brief = sample_ir("brief")
    brief.header.organ_mark = "工作简报"
    assert apply_unit(letter, tpl.unit)[0].header.organ_mark == "示例市卫生健康委员会"
    b, filled = apply_unit(brief, tpl.unit)
    assert b.header.organ_mark == "卫生健康工作简报" and b.meta["brief_issuer"] == "示例市卫生健康委员会办公室" and "简报编印单位" in filled


# ------------------------------------------------------------------ 命令行
def test_template_cli_new_set_validate_show_delete(tmp_path, capsys):
    ws = ["-C", str(tmp_path)]
    assert cli([*ws, "template", "new", "本单位", "--from", "windows-fonts"]) == 0
    assert cli([*ws, "template", "set", "本单位", "elements.title.size=小二", "elements.levels.2=heiti", "unit.organ_mark=示例市卫生健康委员会文件"]) == 0
    out = capsys.readouterr().out
    assert "偏离 GB/T 9704—2012：2 项" in out and "标题字号：二号 → 小二" in out and "二级标题字体：楷体 → 黑体" in out
    assert cli([*ws, "template", "set", "本单位", "elements.title.size=小五"]) == 2
    assert "应为字号名之一" in capsys.readouterr().err
    assert cli([*ws, "template", "validate", "本单位", "--json"]) == 0
    d = json.loads(capsys.readouterr().out)
    assert d["valid"] and sum(c["deviates"] for c in d["changes"]) == 2
    # 手工改坏的模板：validate 报告错误、退出码 2
    p = tmp_path / ".gongwen" / "templates" / "本单位.yaml"
    p.write_text(p.read_text(encoding="utf-8").replace("小二", "特大"), encoding="utf-8")
    assert cli([*ws, "template", "validate", "本单位"]) == 2
    assert "校验不通过" in capsys.readouterr().out
    assert cli([*ws, "template", "list", "--json"]) == 0
    rows = json.loads(capsys.readouterr().out)
    assert any(r["name"] == "本单位" and "error" in r for r in rows)
    assert cli([*ws, "template", "delete", "本单位"]) == 2  # 非交互环境须 --yes
    assert cli([*ws, "template", "delete", "本单位", "--yes"]) == 0 and not p.exists()


def test_task_creation_validates_template_first(tmp_path, capsys):
    ws = ["-C", str(tmp_path), "-c", "layout.render_check=false"]
    assert cli([*ws, "task", "new", "写一份通知", "--template", "不存在"]) == 2
    assert "未找到模板" in capsys.readouterr().err
    assert make_engine(tmp_path).store.list_tasks() == []  # 不留下空任务
    assert cli([*ws, "task", "new", "写一份通知", "--template", "windows-fonts", "--json"]) == 0
    tid = json.loads(capsys.readouterr().out)["task_id"]
    assert make_engine(tmp_path).load_state(tid).options["layout_template"] == "windows-fonts"


def test_task_layout_uses_task_template(tmp_path):
    """任务按创建时指定的模板排版：DOCX 标题字号、报告中的模板与偏离。"""
    from helpers import make_docx, make_xlsx
    from test_engine_e2e import BUDGET, SITUATION

    eng = make_engine(tmp_path)
    TemplateStore(eng.rt.data_dir, tmp_path).save(Template.from_dict({"name": "本单位", "elements": {"title": {"size": "小二"}}}))
    user = default_user()
    st = eng.create_task("写一份向主管部门申请基层医疗示范点建设经费的报告", by=user, hints={"recipients": "示例市人民政府", "issuer_type": "政府部门"}, options={"layout_template": "本单位"})
    eng.add_material(st.task_id, "情况说明.docx", make_docx(SITUATION), by=user, declared=Clearance.PUBLIC)
    eng.add_material(st.task_id, "经费测算表.xlsx", make_xlsx(BUDGET, sheet="经费测算"), by=user, declared=Clearance.PUBLIC)
    st = run_to_review(eng, user, st.task_id)
    from gongwen.schemas.layout import LayoutReport

    rep = eng.store.load_model(st.task_id, "layout_report", LayoutReport)
    assert rep.template == "本单位" and rep.deviations == ["标题字号：二号 → 小二（GB/T 9704—2012 7.3.1，一般）"]
    by = {c.rule_id: c for c in rep.checks}
    assert by["TEMPLATE"].status == "warn" and by["DOCX-TITLE"].status == "pass" and "18磅" in by["DOCX-TITLE"].expected
    assert by["DOCX-TITLE"].actual.startswith("18磅")  # 回读 DOCX：标题按模板为小二（18 磅）


def test_format_without_template_is_unchanged_and_with_template_honours_it(tmp_path, capsys):
    src = tmp_path / "稿.txt"
    src.write_text("关于加强基层医疗示范点建设的通知\n各区卫生健康局：\n为提升基层医疗卫生服务能力，现就有关事项通知如下。\n一、工作目标\n全市新建示范点12个。\n示例市卫生健康委员会\n2026年3月1日\n", encoding="utf-8")
    ws = ["-C", str(tmp_path)]
    assert cli([*ws, "format", str(src), "-o", str(tmp_path / "a"), "--no-render", "--json"]) == 0
    plain = json.loads(capsys.readouterr().out)
    assert plain["template"] == "" and not any(c["rule_id"] in ("TEMPLATE", "TPL-DEV") for c in plain["checks"])
    TemplateStore(tmp_path / ".gongwen", tmp_path).save(
        Template.from_dict({"name": "本单位", "elements": {"title": {"size": "小二"}, "levels": {1: "kaiti"}}, "unit": {"organ_mark": "示例市卫生健康委员会文件", "printer": "示例市卫生健康委员会办公室"}})
    )
    assert cli([*ws, "format", str(src), "-o", str(tmp_path / "b"), "--no-render", "--template", "本单位", "--json"]) == 0
    rep = json.loads(capsys.readouterr().out)
    assert rep["template"] == "本单位" and len(rep["deviations"]) == 2
    by = {c["rule_id"]: c for c in rep["checks"]}
    assert by["DOCX-TITLE"]["status"] == "pass" and "模板“本单位”" in by["DOCX-TITLE"]["expected"]
    assert sum(1 for c in rep["checks"] if c["rule_id"] == "TPL-DEV") == 2
    assert any("单位信息填入：发文机关标志、印发机关" in n for n in rep["render"]["notes"])
    doc = Document(str(tmp_path / "b" / "稿.docx"))
    texts = [p.text for p in doc.paragraphs]
    assert "示例市卫生健康委员会文件" in texts
    h1 = next(p for p in doc.paragraphs if p.text == "一、工作目标").runs[0]
    assert h1._element.rPr.rFonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia") == "楷体_GB2312"
    cells = " ".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "示例市卫生健康委员会办公室" in cells


# ------------------------------------------------------------------ Word 模板
def test_export_dotx_has_template_content_type_and_named_styles(tmp_path, capsys):
    store = TemplateStore(tmp_path / ".gongwen", tmp_path)
    store.save(Template.from_dict({"name": "本单位", "elements": {"levels": {2: "heiti"}}, "unit": {"organ_mark": "示例市卫生健康委员会文件", "doc_number_prefix": "示卫发"}}))
    assert cli(["-C", str(tmp_path), "template", "export-dotx", "本单位", "-o", str(tmp_path / "公文.dotx")]) == 0
    z = zipfile.ZipFile(tmp_path / "公文.dotx")
    ct = z.read("[Content_Types].xml").decode("utf-8")
    assert TEMPLATE_CT in ct and DOCUMENT_CT not in ct
    styles = z.read("word/styles.xml").decode("utf-8")
    for name in ("公文标题", "主送机关", "正文", "一级标题（黑体）", "二级标题（黑体）", "三级标题", "四级标题", "附件说明", "署名", "成文日期", "附注", "版记"):
        assert f'w:val="{name}"' in styles, name
    assert 'w:styleId="GwBody"' in styles and 'w:outlineLvl w:val="0"' in styles
    doc_xml = z.read("word/document.xml").decode("utf-8")
    assert "示例市卫生健康委员会文件" in doc_xml and "示卫发" in doc_xml and 'w:linePitch="579"' in doc_xml
    assert "<w:evenAndOddHeaders/>" in z.read("word/settings.xml").decode("utf-8")
    footers = [n for n in z.namelist() if n.startswith("word/footer")]
    assert len(footers) >= 2 and all("PAGE" in z.read(n).decode("utf-8") for n in footers)  # 7.5：单双页页码
    assert cli(["-C", str(tmp_path), "template", "export-dotx", "本单位", "-o", str(tmp_path / "x.docx")]) == 2


@needs_render
def test_libreoffice_opens_exported_dotx(tmp_path):
    from gongwen.layout.render_check import pdf_pages, to_pdf

    path = export_dotx(LayoutProfile.load(), tmp_path / "国标.dotx")
    pdf = to_pdf(path, tmp_path)
    assert pdf is not None and pdf.is_file()
    text = "".join(l.text for _, _, ls in pdf_pages(pdf) for l in ls)
    assert "【公文标题】" in text and "【发文机关标志】" in text and "—1—" in text.replace(" ", "")


# ------------------------------------------------------------------ 按模板渲染核验
@needs_render
def test_render_checks_compare_against_template_values_and_list_deviations(tmp_path):
    tpl = Template.from_dict({"name": "本单位", "elements": {"title": {"size": "小二"}, "organ_mark": {"top_from_type_area_mm": 30}}})
    ir = sample_ir()
    rep = layout_document(ir, tmp_path / "t", profile=tpl.profile())
    by = {c.rule_id: c for c in rep.checks}
    assert by["LAY-MARK"].expected == "30mm" and by["LAY-MARK"].status == "pass", by["LAY-MARK"].actual
    assert abs(float(by["LAY-MARK"].actual.split("mm")[0]) - 30) <= 1
    assert by["DOCX-TITLE"].status == "pass"  # 有意的小二标题不判为不符合
    devs = [c for c in rep.checks if c.rule_id == "TPL-DEV"]
    assert {c.item for c in devs} == {"模板偏离：标题字号", "模板偏离：发文机关标志上边缘至版心上边缘"}
    mark = next(c for c in devs if "发文机关标志" in c.item)
    assert mark.expected == "35mm（GB/T 9704—2012 7.2.4，规定）" and not mark.conditional and "规定" in mark.note
    assert by["TEMPLATE"].status == "warn" and "偏离国标 2 项（规定 1 项、一般 1 项）" in by["TEMPLATE"].actual
    assert any("按模板“本单位”排版" in n for n in rep.render.notes)
    # 同一份 DOCX 按国标默认参数核验：标志位置不符合 35mm
    from gongwen.layout.docx_compiler import compile_docx
    from gongwen.layout.render_check import check_rendering

    path, fonts = compile_docx(ir, tmp_path / "x" / "x.docx", tpl.profile())
    _, checks = check_rendering(path, ir, LayoutProfile.load(), fonts, tmp_path / "x")
    assert next(c for c in checks if c.rule_id == "LAY-MARK").status == "warn"
