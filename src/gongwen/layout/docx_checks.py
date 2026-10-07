"""生成参数核验：回读 DOCX，核对页面、页边距、字体字号、行距等设置（不等于实际渲染）。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from ..schemas.layout import LayoutCheck
from .docx_compiler import size_label
from .profile import MM_PER_PT, LayoutProfile


def _mm(length) -> float:
    return round(length / 36000, 2) if length is not None else 0.0


def check_docx(path: Path, profile: LayoutProfile, title: str, first_body: str = "") -> list[LayoutCheck]:
    doc = Document(str(path))
    s = doc.sections[0]
    out: list[LayoutCheck] = []
    m = profile.margins
    pg = profile.data["page"]
    out.append(LayoutCheck(rule_id="DOCX-PAGE", item="纸张", expected=f"{pg['width_mm']}×{pg['height_mm']}mm", actual=f"{_mm(s.page_width)}×{_mm(s.page_height)}mm", status="pass" if abs(_mm(s.page_width) - pg["width_mm"]) < 0.5 and abs(_mm(s.page_height) - pg["height_mm"]) < 0.5 else "fail", clause="GB/T 9704—2012 5.1"))
    tol = profile.data["margins"]["top_tol_mm"]
    out.append(
        LayoutCheck(
            rule_id="DOCX-MARGIN-TOP",
            item="天头（上白边）设置",
            expected=f"{profile.data['margins']['top_mm']}mm±{tol}mm" + ("（补偿口径 34.58mm，实务）" if profile.margin_mode == "compensated" else ""),
            actual=f"{_mm(s.top_margin)}mm",
            status="pass" if abs(_mm(s.top_margin) - m["top"]) <= 0.05 else "fail",
            clause="GB/T 9704—2012 5.2.1",
            note="软件页边距量到字框上沿；国标天头量到汉字上沿，二者口径不同" if profile.margin_mode == "compensated" else "",
        )
    )
    out.append(LayoutCheck(rule_id="DOCX-MARGIN-LEFT", item="订口（左白边）设置", expected=f"{profile.data['margins']['left_mm']}mm±{profile.data['margins']['left_tol_mm']}mm", actual=f"{_mm(s.left_margin)}mm", status="pass" if abs(_mm(s.left_margin) - m["left"]) <= profile.data["margins"]["left_tol_mm"] else "fail", clause="GB/T 9704—2012 5.2.1"))
    width = round(_mm(s.page_width) - _mm(s.left_margin) - _mm(s.right_margin), 1)
    height = round(_mm(s.page_height) - _mm(s.top_margin) - _mm(s.bottom_margin), 1)
    ta = profile.data["type_area"]
    out.append(LayoutCheck(rule_id="DOCX-TYPEAREA", item="版心尺寸", expected=f"{ta['width_mm']}×{ta['height_mm']}mm", actual=f"{width}×{height}mm", status="pass" if abs(width - ta["width_mm"]) <= 0.5 and (abs(height - ta["height_mm"]) <= 0.5 or profile.margin_mode == "compensated") else "fail", clause="GB/T 9704—2012 5.2.1"))
    # 正文行距：225mm ÷ 22 行（实务推导值）
    probe = first_body[:8]
    body = [p for p in doc.paragraphs if probe and p.text.startswith(probe)] or [p for p in doc.paragraphs if p.text and p.paragraph_format.first_line_indent and p.paragraph_format.first_line_indent.pt > 20]
    if body:
        ls = body[0].paragraph_format.line_spacing
        pt = ls.pt if hasattr(ls, "pt") else 0
        out.append(LayoutCheck(rule_id="DOCX-LINE", item="正文行距", expected=f"固定值约{profile.line_pt}磅（每面{profile.data['grid']['lines_per_page']}行撑满版心，实务推导）", actual=f"{pt:.2f}磅", status="pass" if abs(pt - profile.line_pt) < 0.1 else "warn", clause="GB/T 9704—2012 5.2.3", level="实务", conditional=True))
    # 标题字体字号：取连续几段拼起来恰好等于标题的段落（标题可能回行为多段）；
    # 不能用“段落文字包含于标题”判断——函的发文机关标志常是标题的开头，会量成红色机关标志
    texts = [p.text for p in doc.paragraphs]
    tparas = []
    for i, t in enumerate(texts):
        acc, j = t, i
        while t and title.startswith(acc) and acc != title and j + 1 < len(texts) and texts[j + 1]:
            j += 1
            acc += texts[j]
        if t and acc == title:
            tparas = doc.paragraphs[i : j + 1]
            break
    # 字体字号与缩进按配置档（模板）的生效参数核对：模板有意改动的项不判为不符合，偏离国标另由模板核验项列出
    tpl = f"（模板“{profile.template}”）" if profile.template else ""
    if tparas:
        r = tparas[0].runs[0]
        rf = r._element.rPr.rFonts.get(qn("w:eastAsia")) if r._element.rPr is not None and r._element.rPr.rFonts is not None else ""
        size = r.font.size.pt if r.font.size else 0
        tel = profile.el("title")
        t_font = tel.get("font", "xiaobiaosong")
        t_pt = profile.size(tel["size"])
        exp_font = profile.font(t_font)
        out.append(LayoutCheck(rule_id="DOCX-TITLE", item="标题字体字号", expected=f"{size_label(tel['size'])}（{t_pt:g}磅）{profile.font_category(t_font)}{tpl}", actual=f"{size:g}磅 {rf}", status="pass" if abs(size - t_pt) < 0.1 and rf == exp_font else "warn", clause="GB/T 9704—2012 7.3.1", conditional=True))
    if body:
        r = body[0].runs[0]
        rf = r._element.rPr.rFonts.get(qn("w:eastAsia")) if r._element.rPr is not None and r._element.rPr.rFonts is not None else ""
        size = r.font.size.pt if r.font.size else 0
        bel = profile.el("body")
        b_font = bel.get("font", "fangsong")
        out.append(LayoutCheck(rule_id="DOCX-BODY", item="正文字体字号", expected=f"{size_label(bel['size'])}（{profile.body_pt:g}磅）{profile.font_category(b_font)}{tpl}", actual=f"{size:g}磅 {rf}", status="pass" if abs(size - profile.body_pt) < 0.1 else "warn", clause="GB/T 9704—2012 7.3.3", conditional=True))
        ind = body[0].paragraph_format.first_line_indent
        chars = (ind.pt / profile.char_pitch_pt) if ind is not None else 0
        want = float(bel.get("first_indent_chars", 2))
        out.append(LayoutCheck(rule_id="DOCX-INDENT", item="自然段左空二字", expected=f"{want:g}字{tpl}", actual=f"{chars:.1f}字", status="pass" if abs(chars - want) < 0.1 else "fail", clause="GB/T 9704—2012 7.3.3"))
    settings = doc.settings.element
    eo = settings.find(qn("w:evenAndOddHeaders")) is not None
    out.append(LayoutCheck(rule_id="DOCX-PAGENO", item="单双页页码位置分设", expected="单页居右、双页居左", actual="已分设" if eo else "未分设", status="pass" if eo else "fail", clause="GB/T 9704—2012 7.5"))
    return out
