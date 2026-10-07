"""DocumentIR → DOCX 编译器（GB/T 9704—2012 通用格式、信函格式、命令（令）格式、纪要格式；另有简报报头与
不设版头的事务文书，均属实务）。

实现要点：
* “空N字”按字距换算为缩进磅值，“空N行”用与正文同高的空段落，不用空格字符填充；
* 正文固定行距 579 缇（28.95 磅：225mm ÷ 22 行约 28.99 磅，向下取整到缇，22 行才不超出版心；实务推导），
  3 号字字距 −0.25 磅使每行 28 字不超出版心；
* 信函格式的发文机关标志（上边缘距上页边 30mm，在天头内）与两条红色双线用锚定于页面的图文框（framePr）定位，
  正文从图文框下方接排；首页另设空页脚，不显示页码（10.1）；
* 版记置于锚定在版心底部的浮动表格中，使末条分隔线与最后一面版心下边缘重合；版记估计高于一面版心时
  改为紧接正文的普通表格（浮动表格会越出版心）；
* 页码奇偶页分设：单页码居右空一字，双页码居左空一字；
* 待补内容（【待……】）以黄色底纹突出，送审时一目了然；
* 文档属性与页眉标注“智能体辅助起草，须人工审核”（《政务领域人工智能大模型部署应用指引》要求做好输出内容标识）。
"""

from __future__ import annotations

import math
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Twips

from ..schemas.ir import Block, DocumentIR
from .profile import MM_PER_PT, LayoutProfile, split_balanced, split_title, text_width_chars

PPR_SEQ = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers",
    "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
    "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
    "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
    "rPr", "sectPr", "pPrChange",
]
RPR_SEQ = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow", "emboss",
    "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs",
    "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
    "specVanish", "oMath",
]


def insert_ordered(parent, el, seq: list[str]) -> None:
    """按 OOXML 架构顺序插入子元素（Word 对元素顺序敏感）。"""
    name = el.tag.split("}")[1]
    idx = seq.index(name)
    for i, child in enumerate(list(parent)):
        cname = child.tag.split("}")[1]
        if cname in seq and seq.index(cname) > idx:
            parent.insert(i, el)
            return
    parent.append(el)


RED = RGBColor(0xFF, 0x00, 0x00)
GRAY = RGBColor(0x80, 0x80, 0x80)
PLACEHOLDER_RE = re.compile(r"(【待[^】]*】)")

# 字位校准（LibreOffice 实测）。固定行距时 LibreOffice 把基线放在行顶之下 0.8 个行高处，与字体无关；汉字字形上缘约在
# 基线之上 0.86 个字高（文泉驿正黑实测 0.855、Noto Serif CJK 0.865），也基本与字体无关。PDF 文本框（字框）的上缘则是
# 基线减去字体的上伸度量（文泉驿正黑 0.963、Noto CJK 1.151），因字体而异——按字框校准会使版面随本机字体而变，
# 因此下列常数都按字形（基线）校准，生成的 DOCX 与本机装了哪种字体无关；渲染核验也按栅格化后的字形测量（render_check.ink_extent）。
# * 图文框内的发文机关标志行距取 1.02 个字高：框顶在标志上边缘之下 0.044（= 0.86 − 0.8×1.02）个字高，字形下缘在框顶下约
#   0.945 个字高（文泉驿正黑 0.975、Noto 0.914 的折中）；
# * 流式排列的标志（mark_line）：补入间距 0.017 个字高（流式排列另有约 0.03 个字高的下移，实测）；
# * 3 号字行距 579 缇时，字形上缘在行顶（框顶）下约 BODY_FRAME_TEXT_TOP_MM（= 0.8×28.95 − 0.86×16 磅）；
# * 页码：按一字线字形的中线校准（FOOTER_PAD_MM；原按字框中线为 0.8mm，一字线实测比字框中线低约 0.6mm）。
# 2026-10 以前按文泉驿正黑的字框上缘校准（0.147、0.13、2.74mm），字形实际比要求低约 1.3～1.5mm；改用 Noto 等开源替代字体
# 渲染后字框上移 2.3～2.4mm 而字形位置不变，故改为按字形校准。
# 实测（2026-10，LibreOffice 24.2，字形上缘／一字线中线；前后两值分别为 Noto 替代字体与文泉驿正黑）：标志 34.9～35.1mm、
# 信函标志 30.0～30.1mm、命令标志 19.8～19.9mm、信函发文字号距双线 4.8～5.1mm（要求 4.9mm）、双线距标志 4.4～3.7mm（要求 4mm）、
# 页码一字线 6.9～7.0mm（要求 7mm）。
MARK_FRAME_ASCENT = 0.044
MARK_FRAME_INK_BOTTOM = 0.945
BODY_FRAME_TEXT_TOP_MM = 3.32
MARK_FLOW_ASCENT = 0.017
FOOTER_PAD_MM = 0.3

_HAO = {"初号": "初号", "小初": "小初", "一号": "1号", "小一": "小1号", "二号": "2号", "小二": "小2号", "三号": "3号", "小三": "小3号", "四号": "4号", "小四": "小4号"}


def size_label(name: str) -> str:
    """字号名的阿拉伯数字写法（“二号”→“2号”），用于核验说明。"""
    return _HAO.get(name, name)


class Compiler:
    def __init__(self, ir: DocumentIR, profile: LayoutProfile, draft_label: bool = True, tight: int = 0):
        self.ir = ir
        self.p = profile
        self.doc = Document()
        self.draft_label = draft_label
        # 署名页无正文时的调整级别（7.3.5.5：所剩空白容不下署名、成文日期时调整行距、字距）：
        # 1 = 正文与附件说明、署名之间的空行减为半行高；2 = 另使正文末段与署名同页
        self.tight = tight
        self.imported = ir.meta.get("source") == "imported"  # 外部文稿：只排版，内容未经本系统核验
        self.fonts_used: set[str] = set()

    # ------------------------------------------------------------------ 底层工具
    @property
    def body_font(self) -> str:
        return self.p.el("body").get("font", "fangsong")

    @property
    def body_size(self) -> str:
        return self.p.el("body")["size"]

    def font_name(self, key: str) -> str:
        name = self.p.font(key) if key in self.p.data["fonts"] else key
        self.fonts_used.add(name)
        return name

    def set_run(self, run, font: str | None = None, size: str | float | None = None, color: RGBColor | None = None, bold: bool = False, spacing: bool = True, highlight: bool = False) -> None:
        font = font or self.body_font
        size = size or self.body_size
        pt = self.p.size(size) if isinstance(size, str) else float(size)
        name = self.font_name(font)
        run.font.size = Pt(pt)
        run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        for attr in ("w:eastAsia", "w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), name)
        rfonts.set(qn("w:hint"), "eastAsia")
        if spacing and abs(pt - self.p.body_pt) < 0.01:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(self.p.data["grid"]["char_spacing_twips"]))
            insert_ordered(rpr, sp, RPR_SEQ)
        if highlight:
            hl = OxmlElement("w:highlight")
            hl.set(qn("w:val"), "yellow")
            insert_ordered(rpr, hl, RPR_SEQ)

    def zi(self, n: float, size: str | None = None) -> Pt:
        """n 个字宽（国标 3.1：一字指一个汉字宽度的距离）。"""
        size = size or self.body_size
        pitch = self.p.char_pitch_pt if size == self.p.el("body")["size"] else self.p.size(size)
        return Pt(n * pitch)

    def para(self, text: str = "", font: str | None = None, size: str | None = None, align=None, first: float = 0, left: float = 0, right: float = 0, line_pt: float | None = None, color=None, bold=False, keep_next=False):
        par = self.doc.add_paragraph()
        self.no_grid(par)
        pf = par.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        # 按缇写入（OOXML 行距单位）：正文 28.95 磅 = 579 缇，22 行 12738 缇不超出 225mm 版心（12756 缇）
        pf.line_spacing = Twips(round((line_pt or self.p.line_pt) * 20))
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.widow_control = False
        if keep_next:
            pf.keep_with_next = True
        if align is not None:
            par.alignment = align
        if first:
            pf.first_line_indent = self.zi(first, size)
        if left:
            pf.left_indent = self.zi(left, size)
        if right:
            pf.right_indent = self.zi(right, size)
        if text:
            self.add_text(par, text, font, size, color, bold)
        return par

    @staticmethod
    def no_grid(par) -> None:
        """段落不对齐文档网格：行高完全由固定行距决定（避免网格把大字号行撑成多格）。"""
        ppr = par._p.get_or_add_pPr()
        if ppr.find(qn("w:snapToGrid")) is None:
            el = OxmlElement("w:snapToGrid")
            el.set(qn("w:val"), "0")
            insert_ordered(ppr, el, PPR_SEQ)

    @staticmethod
    def short_placeholder(text: str, short: str) -> str:
        return short if text.startswith("【待") and len(text) > len(short) else text

    def add_text(self, par, text: str, font: str | None = None, size: str | None = None, color=None, bold=False) -> None:
        for part in PLACEHOLDER_RE.split(text):
            if not part:
                continue
            run = par.add_run(part)
            self.set_run(run, font, size, color, bold, highlight=bool(PLACEHOLDER_RE.fullmatch(part)))

    def blank(self, n: int = 1, line_pt: float | None = None, keep_next: bool = False) -> None:
        for _ in range(n):
            self.para(line_pt=line_pt, keep_next=keep_next)

    def gap_pt(self) -> float | None:
        """正文与附件说明、署名之间空行的行高：调整时减为半行（7.3.5.5 允许调整行距）。"""
        return self.p.line_pt / 2 if self.tight >= 1 else None

    def keep_last_with_next(self) -> None:
        """使当前最后一个段落与下一段同页（最后一个元素是表格时不处理）。"""
        last = next((el for el in reversed(self.doc.element.body) if el.tag != qn("w:sectPr")), None)
        if last is not None and last.tag == qn("w:p"):
            ppr = last.get_or_add_pPr()
            if ppr.find(qn("w:keepNext")) is None:
                insert_ordered(ppr, OxmlElement("w:keepNext"), PPR_SEQ)

    @staticmethod
    def border(par, edge: str, color: str = "000000", width_pt: float = 0.75, space_pt: float = 0) -> None:
        ppr = par._p.get_or_add_pPr()
        bdr = ppr.find(qn("w:pBdr"))
        if bdr is None:
            bdr = OxmlElement("w:pBdr")
            insert_ordered(ppr, bdr, PPR_SEQ)
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(max(2, int(round(width_pt * 8)))))
        el.set(qn("w:space"), str(int(round(space_pt))))
        el.set(qn("w:color"), color)
        order = ["top", "left", "bottom", "right", "between", "bar"]
        for i, child in enumerate(list(bdr)):
            if order.index(child.tag.split("}")[1]) > order.index(edge):
                bdr.insert(i, el)
                break
        else:
            bdr.append(el)

    @staticmethod
    def frame(par, y_mm: float, width_mm: float) -> None:
        """把段落放入锚定于页面的图文框：框顶距上页边 y_mm，宽 width_mm，以版心为准水平居中。

        正文不在图文框两侧环绕（notBeside），遇到图文框时从其下方接排；相邻段落的图文框设置相同则合为一个框。"""
        ppr = par._p.get_or_add_pPr()
        fp = OxmlElement("w:framePr")
        for k, v in (("w:w", round(width_mm / MM_PER_PT * 20)), ("w:hSpace", 0), ("w:vSpace", 0), ("w:wrap", "notBeside"), ("w:vAnchor", "page"), ("w:hAnchor", "margin"), ("w:xAlign", "center"), ("w:y", round(y_mm / MM_PER_PT * 20))):
            fp.set(qn(k), str(v))
        insert_ordered(ppr, fp, PPR_SEQ)

    def double_rule(self, y_mm: float, top_pt: float, bottom_pt: float) -> float:
        """红色双线（信函格式 10.1）：只有上下边框的图文框段落，上缘距上页边 y_mm，两线间隔为段落行高。
        返回双线总高（mm）。用两条单线边框而不用 OOXML 的双线边框样式：后者粗细线的上下次序各软件解释不一。"""
        lt = self.p.data["letter"]
        par = self.para(line_pt=lt["rule_gap_pt"])
        self.border(par, "top", "FF0000", top_pt)
        self.border(par, "bottom", "FF0000", bottom_pt)
        self.frame(par, y_mm, lt["rule_length_mm"])
        return (top_pt + lt["rule_gap_pt"] + bottom_pt) * MM_PER_PT

    def field(self, par, instr: str, font: str, size: str) -> None:
        run = par.add_run()
        self.set_run(run, font, size, spacing=False)
        b = OxmlElement("w:fldChar")
        b.set(qn("w:fldCharType"), "begin")
        t = OxmlElement("w:instrText")
        t.set(qn("xml:space"), "preserve")
        t.text = instr
        s = OxmlElement("w:fldChar")
        s.set(qn("w:fldCharType"), "separate")
        r2 = OxmlElement("w:t")
        r2.text = "1"
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), "end")
        run._r.append(b)
        run._r.append(t)
        run._r.append(s)
        run._r.append(r2)
        run._r.append(e)

    # ------------------------------------------------------------------ 版面
    def setup_page(self) -> None:
        s = self.doc.sections[0]
        page = self.p.data["page"]
        m = self.p.margins
        s.page_width, s.page_height = Mm(page["width_mm"]), Mm(page["height_mm"])
        s.top_margin, s.bottom_margin = Mm(m["top"]), Mm(m["bottom"])
        s.left_margin, s.right_margin = Mm(m["left"]), Mm(m["right"])
        pn = self.p.el("page_number")
        num_mm = self.p.size(pn["size"]) * MM_PER_PT
        # 一字线上距版心下边缘 7mm（推算页脚距离：页高 − 版心下边缘 − 7mm − 半个字高）；FOOTER_PAD_MM 为字框留白的实测校准
        footer_mm = page["height_mm"] - (self.p.data["margins"]["top_mm"] + self.p.data["type_area"]["height_mm"]) - pn["dash_offset_mm"] - num_mm * 0.5 - FOOTER_PAD_MM
        s.footer_distance = Mm(max(5.0, footer_mm))
        s.header_distance = Mm(12)
        sect = s._sectPr
        # 只对齐行网格；每行 28 字通过 3 号字字距实现（字符网格会使大字号标题被拉开）
        for old in sect.findall(qn("w:docGrid")):
            sect.remove(old)
        grid = OxmlElement("w:docGrid")
        grid.set(qn("w:type"), "lines")
        grid.set(qn("w:linePitch"), str(int(round(self.p.line_pt * 20))))
        sect.append(grid)
        # 默认样式：正文字体字号（国标一般为 3 号仿宋）
        st = self.doc.styles["Normal"]
        st.font.size = Pt(self.p.body_pt)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        for attr in ("w:eastAsia", "w:ascii", "w:hAnsi"):
            rf.set(qn(attr), self.font_name(self.body_font))
        # 奇偶页不同页脚
        settings = self.doc.settings.element
        if settings.find(qn("w:evenAndOddHeaders")) is None:
            settings.append(OxmlElement("w:evenAndOddHeaders"))

    def page_numbers(self) -> None:
        s = self.doc.sections[0]
        pn = self.p.el("page_number")
        size = pn["size"]
        font = pn["font"]
        indent = Pt(pn["indent_chars"] * self.p.size(size))
        dash = pn.get("style", "dash") != "plain"  # 模板可改为不加一字线（偏离 7.5，模板偏离中列出）
        for footer, align in ((s.footer, WD_ALIGN_PARAGRAPH.RIGHT), (s.even_page_footer, WD_ALIGN_PARAGRAPH.LEFT)):
            par = footer.paragraphs[0]
            par.alignment = align
            if align == WD_ALIGN_PARAGRAPH.RIGHT:
                par.paragraph_format.right_indent = indent
            else:
                par.paragraph_format.left_indent = indent
            if dash:
                r = par.add_run("— ")
                self.set_run(r, font, size, spacing=False)
            self.field(par, "PAGE", font, size)
            if dash:
                r = par.add_run(" —")
                self.set_run(r, font, size, spacing=False)
        no_first_number = self.ir.format_type == "letter" and not self.p.data["letter"].get("first_page_number", True)
        if no_first_number:
            s.different_first_page_header_footer = True
        if self.draft_label:
            key = "format_label" if self.imported else "draft_label"
            label = self.p.data.get(key, self.p.data.get("draft_label", "")).format(status=self.ir.status.value)
            for header in (s.header, s.even_page_header):
                hp = header.paragraphs[0]
                hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = hp.add_run(label)
                self.set_run(r, "songti", "小四", color=GRAY, spacing=False)
            if self.ir.format_type == "letter":
                hp = s.first_page_header.paragraphs[0]
                hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = hp.add_run(label)
                self.set_run(r, "songti", "小四", color=GRAY, spacing=False)
        if no_first_number:
            # 首页不显示页码（10.1）：另设空的首页页脚。须在首页页眉之后添加——首页页脚的引用排在首页页眉之前时，
            # LibreOffice 首页仍显示奇数页页脚（实测）
            s.first_page_footer.is_linked_to_previous = False

    # ------------------------------------------------------------------ 版头
    def mark_size(self, mark: str, max_size: str | None = None) -> float:
        """发文机关标志字号：不大于 max_size（默认小初），且一行排得下（国标未规定字号，实务按版面确定）。"""
        cap = self.p.size(max_size or self.p.el("organ_mark")["max_size"])
        return min(cap, (self.p.type_width_pt * 0.96) / max(1.0, text_width_chars(mark or "文")))

    @property
    def mark_font(self) -> str:
        return self.p.el("organ_mark").get("font", "xiaobiaosong")

    @property
    def mark_color(self) -> RGBColor:
        return RGBColor.from_string(str(self.p.el("organ_mark").get("color", "FF0000")).lstrip("#").upper())

    def mark_line(self, mark: str, top_mm: float, size_pt: float, lines_used: int = 0) -> None:
        """红色小标宋的发文机关标志（简报名称）居中，上边缘至版心上边缘 top_mm；其上已排 lines_used 行。"""
        # 按字形校准（见 MARK_FLOW_ASCENT）：补入间距，使字形上缘落在 top_mm 处
        gap_pt = top_mm / MM_PER_PT - lines_used * self.p.line_pt + (MARK_FLOW_ASCENT * size_pt if mark else 0)
        if gap_pt > 1:
            self.para(line_pt=gap_pt)
        if mark:
            par = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, line_pt=size_pt * 1.02)
            run = par.add_run(mark)
            self.set_run(run, self.mark_font, size_pt, color=self.mark_color, spacing=False)

    def header_block(self) -> None:
        ir, h = self.ir, self.ir.header
        fmt = ir.format_type
        if fmt == "plain":
            return  # 事务文书（方案、总结、讲话稿等）不设版头
        if fmt == "letter":
            self.letter_header()
            return
        if fmt == "brief":
            self.brief_header()
            return
        lines_used = 0
        for val, key in ((h.copy_no, "copy_no"), (h.secrecy, "secrecy"), (h.urgency, "urgency")):
            if val:
                el = self.p.el(key)
                self.para(val, font=el["font"], size=el["size"])
                lines_used += 1
        mark = h.organ_mark or ""
        if fmt == "command":
            top_mm = self.p.data["command"]["organ_mark_top_from_type_area_mm"]
        elif fmt == "jiyao":
            top_mm = self.p.data["jiyao"]["mark_top_from_type_area_mm"]
        else:
            top_mm = self.p.el("organ_mark")["top_from_type_area_mm"]
        self.mark_line(mark, top_mm, self.mark_size(mark), lines_used)
        if fmt in ("jiyao",):
            rule = self.para(line_pt=self.p.line_pt)
            rr = self.p.el("red_rule")
            self.border(rule, "bottom", rr.get("color", "FF0000"), rr["width_pt"], 0)
            self.blank(1)
            return
        dn = self.p.el("doc_number")
        dn_font, dn_size = dn.get("font", self.body_font), dn.get("size", self.body_size)
        if fmt == "command":
            # 命令（令）格式（10.2）：发文机关标志下空二行居中编排令号，令号下空二行编排正文；不设分隔线与标题
            n = self.p.data["command"]["number_blank_lines"]
            self.blank(n)
            self.para(self.short_placeholder(h.doc_number, "第【待编号】号"), font=dn_font, size=dn_size, align=WD_ALIGN_PARAGRAPH.CENTER)
            self.blank(n)
            return
        self.blank(dn["blank_lines_after_mark"])
        upward = ir.direction == "上行文"
        doc_number = self.short_placeholder(h.doc_number, "【待编号】")
        if upward:
            sg = self.p.el("signer")
            par = self.para(first=1)
            self.add_text(par, doc_number, font=dn_font, size=dn_size)
            tab_pos = Pt(self.p.type_width_pt - self.p.char_pitch_pt)
            par.paragraph_format.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.RIGHT)
            r = par.add_run("\t签发人：")
            self.set_run(r, sg.get("label_font", "fangsong"), sg.get("size", self.body_size))
            signers = "　".join(h.signers) if h.signers else "【待签发人】"
            for part in PLACEHOLDER_RE.split(signers):
                if part:
                    rr = par.add_run(part)
                    self.set_run(rr, sg.get("name_font", "kaiti"), sg.get("size", self.body_size), highlight=bool(PLACEHOLDER_RE.fullmatch(part)))
        else:
            par = self.para(align=WD_ALIGN_PARAGRAPH.CENTER)
            self.add_text(par, doc_number, font=dn_font, size=dn_size)
        rr = self.p.el("red_rule")
        self.border(par, "bottom", rr["color"], rr["width_pt"], rr["below_doc_number_mm"] / MM_PER_PT)

    def letter_header(self) -> None:
        """信函格式（10.1）：发文机关标志上边缘至上页边 30mm；标志下 4mm 处印一条红色双线（上粗下细），距下页边
        20mm 处印一条红色双线（上细下粗），线长均为 170mm，居中排布；份号、密级和保密期限、紧急程度顶格居版心左边缘、
        发文字号顶格居版心右边缘，编排在第一条双线下，与该线的距离为 3 号汉字高度的 7/8。

        标志在天头内、第二条双线在版心之下，流式段落排不到，因此各要素都放在锚定于页面的图文框中，由计算得到的
        框顶位置定位；正文从最下一个图文框下方接排。图文框的锚定段落都在首页，第二条双线因而只出现在首页。"""
        h, lt = self.ir.header, self.p.data["letter"]
        mark = h.organ_mark or ""
        y = lt["organ_mark_top_from_page_mm"]  # 标志上边缘
        if mark:
            size_pt = self.mark_size(mark)
            size_mm = size_pt * MM_PER_PT
            par = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, line_pt=size_pt * 1.02)
            run = par.add_run(mark)
            self.set_run(run, self.mark_font, size_pt, color=self.mark_color, spacing=False)
            frame_y = y + MARK_FRAME_ASCENT * size_mm
            self.frame(par, frame_y, self.p.data["type_area"]["width_mm"])
            y = frame_y + MARK_FRAME_INK_BOTTOM * size_mm  # 标志下边缘
        thick, thin = lt["rule_thick_pt"], lt["rule_thin_pt"]
        rule_y = y + lt["rule_below_mark_mm"]
        rule_h = self.double_rule(rule_y, thick, thin)
        self.double_rule(self.p.data["page"]["height_mm"] - lt["bottom_rule_from_page_mm"] - rule_h, thin, thick)
        lefts = [(val, key) for val, key in ((h.copy_no, "copy_no"), (h.secrecy, "secrecy"), (h.urgency, "urgency")) if val]
        doc_number = self.short_placeholder(h.doc_number, "【待编号】") if h.doc_number else ""
        if not (lefts or doc_number):
            return
        # 第一个要素与发文字号同排一行（左、右），其余要素自上而下分行；字框上缘距第一条双线 7/8 个 3 号字高
        dn = self.p.el("doc_number")
        dn_font, dn_size = dn.get("font", self.body_font), dn.get("size", self.body_size)
        el_y = rule_y + rule_h + lt["element_gap_ratio"] * self.p.body_pt * MM_PER_PT - BODY_FRAME_TEXT_TOP_MM
        for i in range(max(1, len(lefts))):
            par = self.para()
            if i < len(lefts):
                val, key = lefts[i]
                el = self.p.el(key)
                self.add_text(par, val, font=el["font"], size=el["size"])
            if i == 0 and doc_number:
                par.paragraph_format.tab_stops.add_tab_stop(Pt(self.p.type_width_pt), WD_TAB_ALIGNMENT.RIGHT)
                self.add_text(par, f"\t{doc_number}", font=dn_font, size=dn_size)
            self.frame(par, el_y, self.p.data["type_area"]["width_mm"])

    def brief_header(self) -> None:
        """简报报头（实务：GB/T 9704—2012 未规定简报格式）：红色简报名称居中，下空一行居中编排期号，再空一行左编印单位、
        右日期，其下 4mm 印红色分隔线；分隔线下空二行编排文章标题。简报不署名、不设版记。"""
        h, b = self.ir.header, self.p.data["brief"]
        mark = h.organ_mark or "工作简报"
        self.mark_line(mark, b["mark_top_from_type_area_mm"], self.mark_size(mark, b["mark_size"]))
        self.blank(b["issue_blank_lines"])
        if h.doc_number:
            self.para(self.short_placeholder(h.doc_number, "第【待编号】期"), align=WD_ALIGN_PARAGRAPH.CENTER)
        self.blank(b["issuer_blank_lines"])
        issuer = self.ir.meta.get("brief_issuer") or next((o for o in self.ir.signature.organs if o), "") or b["issuer_placeholder"]
        date = self.ir.meta.get("brief_date") or b["date_placeholder"]
        par = self.para()
        self.add_text(par, issuer)
        par.paragraph_format.tab_stops.add_tab_stop(Pt(self.p.type_width_pt), WD_TAB_ALIGNMENT.RIGHT)
        self.add_text(par, f"\t{date}")
        self.border(par, "bottom", "FF0000", b["rule_width_pt"], b["rule_below_mm"] / MM_PER_PT)

    # ------------------------------------------------------------------ 主体
    def title_block(self) -> None:
        el = self.p.el("title")
        fmt = self.ir.format_type
        if self.ir.title:
            if fmt == "plain":
                before = 0  # 事务文书不设版头：标题从版心第一行起排（实务）
            elif fmt == "letter":
                before = self.p.data["letter"]["title_blank_lines"]  # 与其上最后一个要素相距二行（10.1）
            else:
                before = el["blank_lines_before"]  # 红色分隔线下空二行（7.3.1）
            self.blank(before)
            issuer = self.ir.signature.organs[0] if self.ir.signature.organs else ""
            for line in split_title(self.ir.title, el["max_chars_per_line"], issuer):
                self.para(line, font=el.get("font", "xiaobiaosong"), size=el["size"], align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        if self.ir.title_note:
            # 题注（实务）：标题下居中，楷体 3 号；一行排不下时均衡回行，不让“通过）”之类的短尾单独成行
            tn = self.p.el("title_note")
            for line in split_balanced(self.ir.title_note, self.p.data["grid"]["chars_per_line"]):
                par = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
                self.add_text(par, line, font=tn["font"], size=tn["size"])
        if self.ir.salutation:
            self.blank(1)
            self.para(self.ir.salutation)

    def recipients_block(self) -> None:
        if not self.ir.recipients:
            return
        self.blank(self.p.el("recipients")["blank_lines_before"])
        self.para("、".join(r.rstrip("：:") for r in self.ir.recipients) + "：")

    def body_block(self, blocks: list[Block]) -> None:
        levels = self.p.el("levels")
        indent = self.p.el("body").get("first_indent_chars", 2)
        for b in blocks:
            if b.kind == "heading":
                font = levels.get(b.level, levels.get(str(b.level), self.body_font))
                # 只有单独成段的层次标题才与下段同页；带正文的列项（如“（一）……。”）若也设，会连成一串把整段推到下一面
                par = self.para(first=indent, keep_next=not (b.inline_heading and b.sentences))
                self.add_text(par, f"{b.label}{b.heading}", font=font)
                if b.inline_heading and b.sentences:
                    self.add_text(par, "".join(s.text for s in b.sentences))
            elif b.kind == "table" and b.table:
                self.table(b.table)
            else:
                text = b.text()
                if text:
                    self.para(text, first=indent)

    def table(self, rows: list[list[str]]) -> None:
        ncols = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=ncols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j in range(ncols):
                cell = t.cell(i, j)
                cell.text = ""
                par = cell.paragraphs[0]
                self.no_grid(par)
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                txt = row[j] if j < len(row) else ""
                for part in PLACEHOLDER_RE.split(txt):
                    if part:
                        r = par.add_run(part)
                        self.set_run(r, "heiti" if i == 0 else "fangsong", "四号", spacing=False, highlight=bool(PLACEHOLDER_RE.fullmatch(part)))

    def attachment_note_block(self) -> None:
        notes = self.ir.attachment_notes
        if not notes:
            return
        el = self.p.el("attachment_note")
        # 附件说明与署名同页；仍放不下时（tight=2）正文末段也与署名同页，避免署名页只有附件说明而无正文（7.3.5.5）
        if self.tight >= 2:
            self.keep_last_with_next()
        self.blank(el["blank_lines_before"], line_pt=self.gap_pt(), keep_next=True)
        base = el["left_indent_chars"]
        for i, n in enumerate(notes):
            label = f"{n.seq}." if len(notes) > 1 else ""
            name = n.name.rstrip("。，；：.,;")
            if i == 0:
                par = self.para(left=base + 3 + (1 if label else 0), first=-(3 + (1 if label else 0)), keep_next=True)
                self.add_text(par, f"附件：{label}{name}")
            else:
                par = self.para(left=base + 3 + 1, first=-1, keep_next=True)
                self.add_text(par, f"{label}{name}")

    def signature_block(self) -> None:
        sig = self.ir.signature
        organs = [o for o in sig.organs if o]
        date = sig.date
        # 署名、成文日期不能脱离正文单独落到下一面：前一段、空行与署名行都与下段同页（成文日期本身不设）
        self.keep_last_with_next()
        if sig.seal_mode == "signature_stamp":
            # 加盖签发人签名章（7.3.5.3）：正文下空二行右空四字加盖签名章，签名章左空二字标注签发人职务；
            # 签名章下空一行右空四字编排成文日期
            el = self.p.el("signature_seal")
            right = el["date_right_indent_chars"]
            self.blank(2, line_pt=self.gap_pt(), keep_next=True)
            self.para(f"{sig.signer_title or '【待补：签发人职务】'}　　【签名章】", align=WD_ALIGN_PARAGRAPH.RIGHT, right=right, keep_next=True)
            self.blank(1, keep_next=True)
            self.para(date, align=WD_ALIGN_PARAGRAPH.RIGHT, right=right)
        elif sig.seal_mode == "none":
            pass  # 简报、讲话稿等不署名
        elif sig.seal_mode == "seal":
            el = self.p.el("signature_seal")
            self.blank(2, line_pt=self.gap_pt(), keep_next=True)
            dw = text_width_chars(date)
            right_date = el["date_right_indent_chars"]
            for o in organs:
                ow = text_width_chars(o)
                right = max(0.0, right_date + (dw - ow) / 2)
                self.para(o, align=WD_ALIGN_PARAGRAPH.RIGHT, right=right, keep_next=True)
            self.para(date, align=WD_ALIGN_PARAGRAPH.RIGHT, right=right_date)
        else:
            el = self.p.el("signature_noseal")
            self.blank(1, line_pt=self.gap_pt(), keep_next=True)
            ow = max((text_width_chars(o) for o in organs), default=0)
            dw = text_width_chars(date)
            organ_right = el["organ_right_indent_chars"]
            date_right = organ_right - el["date_shift_chars"] - (dw - ow)
            if date_right < organ_right - el["date_shift_chars"] and dw > ow:
                date_right = el["organ_right_indent_chars"]
                organ_right = date_right + el["date_shift_chars"] + (dw - ow)
            elif dw <= ow:
                date_right = max(0.0, organ_right + (ow - dw) - el["date_shift_chars"])
            for o in organs:
                self.para(o, align=WD_ALIGN_PARAGRAPH.RIGHT, right=organ_right, keep_next=True)
            self.para(date, align=WD_ALIGN_PARAGRAPH.RIGHT, right=max(0.0, date_right))
        if self.ir.note:
            note = self.ir.note.strip("（）()")
            self.para(f"（{note}）", left=self.p.el("note")["left_indent_chars"])

    def attendees_block(self) -> None:
        if not self.ir.attendees:
            return
        # 与署名同理：出席名单不单独落到没有正文的下一面（tight 调整见 gap_pt）
        if self.tight >= 2:
            self.keep_last_with_next()
        self.blank(1, line_pt=self.gap_pt(), keep_next=True)
        keys = [k for k in ("出席", "请假", "列席") if self.ir.attendees.get(k)]
        for key in keys:
            names = self.ir.attendees.get(key)
            par = self.para(left=2 + 3, first=-3, keep_next=key != keys[-1])
            r = par.add_run(f"{key}：")
            self.set_run(r, self.p.data["jiyao"]["attendee_label_font"], "三号")
            self.add_text(par, "、".join(names))

    def attachments_block(self) -> None:
        el = self.p.el("attachment")
        for att in self.ir.attachments:
            par = self.para()
            par.add_run().add_break(WD_BREAK.PAGE)
            # 只有一个附件时附件说明不编顺序号，附件页标识相应为“附件”（7.3.4、7.3.7）
            label = "附件" if len(self.ir.attachments) == 1 and len(self.ir.attachment_notes) <= 1 else f"附件{att.seq}"
            self.para(label, font=el["label_font"], size=el["size"])
            self.blank(1)
            tel = self.p.el("title")
            for line in split_title(att.title, tel["max_chars_per_line"]):
                self.para(line, font=tel.get("font", "xiaobiaosong"), size=tel["size"], align=WD_ALIGN_PARAGRAPH.CENTER)
            self.blank(1)
            self.body_block(att.blocks)

    def imprint_block(self) -> None:
        """版记：以锚定于版心底部的浮动表格实现，首末条分隔线为粗线、中间为细线（7.4.1 推荐值）。"""
        if self.ir.format_type in ("plain", "brief"):
            return  # 事务文书、简报不是正式行文，不设版记（实务）
        imp = self.ir.imprint
        el = self.p.el("imprint")
        size = el["size"]
        ind = el["indent_chars"]
        rows: list[tuple[str, str]] = []
        for label, items in (("主送", imp.main_moved), ("抄送", imp.cc)):
            if items:
                rows.append((label, f"{label}：{'，'.join(items)}。"))
        if self.ir.format_type != "letter":
            if not (rows or imp.printer):
                return
            date = imp.print_date if imp.print_date.endswith("印发") or imp.print_date.startswith("【") else f"{imp.print_date}印发"
            rows.append(("印发", f"{imp.printer or ''}\t{date}"))
        if not rows:
            return
        outer = max(2, int(round(el["outer_rule_mm"] / MM_PER_PT * 8)))
        inner = max(2, int(round(el["inner_rule_mm"] / MM_PER_PT * 8)))
        letter = self.ir.format_type == "letter"
        char = self.p.size(size)
        # 按行数估计版记高度。末页剩余版面不够时，渲染器会把浮动版记整体移到下一面底部（仍符合 7.4.1）；
        # 但版记高于一面版心（如大量主送机关移入版记）时，锚定在版心底部必然越出版心乃至页面，改为紧接正文的普通表格
        per_line = (self.p.type_width_pt - char * ind * 2) / char
        n_lines = sum(1 if kind == "印发" else 1 + math.ceil(max(0.0, text_width_chars(text) - per_line) / (per_line - 3)) for kind, text in rows)
        floating = n_lines * self.p.line_pt * 0.95 <= self.p.type_height_pt - self.p.line_pt
        if not floating:
            self.blank(1)
        table = self.doc.add_table(rows=len(rows), cols=1)
        tbl = table._tbl
        tblPr = tbl.tblPr
        if floating:
            pos = OxmlElement("w:tblpPr")
            for k, v in (("w:leftFromText", "0"), ("w:rightFromText", "0"), ("w:vertAnchor", "margin"), ("w:horzAnchor", "margin"), ("w:tblpXSpec", "center"), ("w:tblpYSpec", "bottom")):
                pos.set(qn(k), v)
            tblPr.insert(1 if tblPr.find(qn("w:tblStyle")) is not None else 0, pos)
        else:
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
        width = str(int(round(self.p.type_width_pt * 20)))
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:w"), width)
        tblW.set(qn("w:type"), "dxa")
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{edge}")
            if edge in ("top", "bottom") and not letter:
                b.set(qn("w:val"), "single")
                b.set(qn("w:sz"), str(outer))
                b.set(qn("w:color"), "000000")
            else:
                b.set(qn("w:val"), "nil")
            borders.append(b)
        tblPr.append(borders)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        mar = OxmlElement("w:tblCellMar")
        for edge in ("left", "right"):
            m = OxmlElement(f"w:{edge}")
            m.set(qn("w:w"), "0")
            m.set(qn("w:type"), "dxa")
            mar.append(m)
        tblPr.append(mar)
        for i, (kind, text) in enumerate(rows):
            cell = table.cell(i, 0)
            cell.width = Pt(self.p.type_width_pt)
            par = cell.paragraphs[0]
            self.no_grid(par)
            pf = par.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(self.p.line_pt * 0.95)
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.left_indent = Pt(char * ind)
            pf.right_indent = Pt(char * ind)
            if kind == "印发":
                pf.tab_stops.add_tab_stop(Pt(self.p.type_width_pt - char * ind * 2), WD_TAB_ALIGNMENT.RIGHT)
                prev_is_cc = i > 0
                if prev_is_cc and not letter:
                    tcPr = cell._tc.get_or_add_tcPr()
                    tcb = OxmlElement("w:tcBorders")
                    top = OxmlElement("w:top")
                    top.set(qn("w:val"), "single")
                    top.set(qn("w:sz"), str(inner))
                    top.set(qn("w:color"), "000000")
                    tcb.append(top)
                    tcPr.append(tcb)
            else:
                pf.left_indent = Pt(char * (ind + 3))
                pf.first_line_indent = Pt(-char * 3)
            for part in PLACEHOLDER_RE.split(text):
                if part:
                    r = par.add_run(part)
                    self.set_run(r, el.get("font", "fangsong"), size, spacing=False, highlight=bool(PLACEHOLDER_RE.fullmatch(part)))

    # ------------------------------------------------------------------
    def compile(self, path: str | Path) -> Path:
        self.setup_page()
        self.page_numbers()
        self.header_block()
        if self.ir.format_type != "jiyao" or self.ir.title:
            self.title_block()
        self.recipients_block()
        self.body_block(self.ir.blocks)
        self.attachment_note_block()
        if self.ir.format_type == "jiyao":
            self.attendees_block()
        elif self.ir.format_type != "brief":  # 简报不署名（实务）
            self.signature_block()
        self.attachments_block()
        self.imprint_block()
        cp = self.doc.core_properties
        cp.title = self.ir.title[:200]
        cp.subject = f"{self.ir.genre or self.ir.material_type or ''}｜{self.ir.status.value}"
        if self.imported:
            cp.keywords = "公文智能体;排版;内容未经系统核验;须人工审核"
            basis = f"{self.p.id}（模板：{self.p.template}）" if self.p.template else self.p.id
            cp.comments = f"本稿由公文智能体按 {basis} 排版，正文内容来自外部文稿、未经本系统核验，须经人工审核。"
            cp.author = "GONGWEN 公文智能体（排版）"
        else:
            cp.keywords = "公文智能体;AI辅助起草;须人工审核"
            cp.comments = f"doc_id={self.ir.doc_id}; version={self.ir.version}; status={self.ir.status.value}; 本稿由公文智能体辅助起草，须经人工审核。"
            cp.author = "GONGWEN 公文智能体（辅助起草）"
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return path


def compile_docx(ir: DocumentIR, path: str | Path, profile: LayoutProfile | None = None, draft_label: bool = True, tight: int = 0) -> tuple[Path, set[str]]:
    c = Compiler(ir, profile or LayoutProfile.load(), draft_label=draft_label, tight=tight)
    out = c.compile(path)
    return out, c.fonts_used
