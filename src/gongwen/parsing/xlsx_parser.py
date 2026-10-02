"""XLSX 解析：保留工作表与单元格定位（Sheet1!B5）、单元格批注、隐藏工作表/行/列；
对公式单元格，若文件未缓存计算结果，则对 SUM/四则运算等简单公式做本地求值并标注。"""

from __future__ import annotations

import io
import re

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from ..schemas.sources import SourceRelation
from .base import NOTE_RE, HiddenContent, ParseResult, UnitBuilder, table_from_rows

_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_FUNC = re.compile(r"(SUM|AVERAGE|MIN|MAX)\(([^()]*)\)", re.IGNORECASE)


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return f"{v:.6f}".rstrip("0").rstrip(".")
    return str(v).strip()


class _Evaluator:
    """仅支持只读的简单公式：单元格引用、区域 SUM/AVERAGE/MIN/MAX 与四则运算。"""

    def __init__(self, values: dict[str, object], formulas: dict[str, str]):
        self.values = values
        self.formulas = formulas
        self.stack: set[str] = set()

    def cell(self, ref: str) -> float:
        ref = ref.replace("$", "")
        if ref in self.formulas and self.values.get(ref) in (None, ""):
            if ref in self.stack:
                raise ValueError("循环引用")
            self.stack.add(ref)
            try:
                self.values[ref] = self.eval(self.formulas[ref])
            finally:
                self.stack.discard(ref)
        v = self.values.get(ref)
        if v in (None, ""):
            return 0.0
        return float(v)

    def _range(self, rng: str) -> list[float]:
        min_col, min_row, max_col, max_row = range_boundaries(rng.replace("$", ""))
        return [
            self.cell(f"{get_column_letter(c)}{r}")
            for r in range(min_row, max_row + 1)
            for c in range(min_col, max_col + 1)
        ]

    def eval(self, formula: str) -> float:
        expr = formula.lstrip("=").upper()
        if "!" in expr:
            raise ValueError("不支持跨表引用")

        def fn(m: re.Match) -> str:
            name, args = m.group(1).upper(), m.group(2)
            vals: list[float] = []
            for a in args.split(","):
                a = a.strip()
                if ":" in a:
                    vals += self._range(a)
                elif a:
                    vals.append(float(self.eval(a)))
            if name == "SUM":
                return repr(sum(vals))
            if name == "AVERAGE":
                return repr(sum(vals) / len(vals) if vals else 0.0)
            if name == "MIN":
                return repr(min(vals) if vals else 0.0)
            return repr(max(vals) if vals else 0.0)

        prev = None
        while prev != expr:
            prev = expr
            expr = _FUNC.sub(fn, expr)
        expr = _REF.sub(lambda m: repr(self.cell(f"{m.group(1)}{m.group(2)}")), expr)
        if not re.fullmatch(r"[\d\.\s\+\-\*/\(\)eE]+", expr):
            raise ValueError(f"不支持的公式：{formula}")
        return float(eval(expr, {"__builtins__": {}}, {}))  # noqa: S307 - 已限定字符集


def parse_xlsx(material_id: str, data: bytes) -> ParseResult:
    b = UnitBuilder(material_id)
    res = b.result
    wb_f = load_workbook(io.BytesIO(data), data_only=False)
    wb_v = load_workbook(io.BytesIO(data), data_only=True)
    t_no = 0
    for ws_f in wb_f.worksheets:
        ws_v = wb_v[ws_f.title]
        if ws_f.sheet_state != "visible":
            texts = [_fmt(c.value) for row in ws_f.iter_rows() for c in row if c.value not in (None, "")]
            res.hidden.append(HiddenContent("隐藏工作表", f"{ws_f.title}：" + "；".join(texts[:50])))
            continue
        hidden_rows = {i for i, d in ws_f.row_dimensions.items() if d.hidden}
        hidden_cols = {column_index_from_string(k) for k, d in ws_f.column_dimensions.items() if d.hidden}
        values: dict[str, object] = {}
        formulas: dict[str, str] = {}
        for row in ws_f.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas[c.coordinate] = c.value
                    values[c.coordinate] = ws_v[c.coordinate].value
                else:
                    values[c.coordinate] = c.value
        ev = _Evaluator(dict(values), formulas)
        grid: list[list[str]] = []
        t_no += 1
        tid = f"{material_id}.t{t_no}"
        for row in ws_f.iter_rows():
            out_row = []
            for c in row:
                coord = c.coordinate
                attrs = {"sheet": ws_f.title, "table": tid}
                val = values.get(coord)
                if coord in formulas:
                    attrs["formula"] = formulas[coord]
                    if val in (None, ""):
                        try:
                            val = ev.eval(formulas[coord])
                            attrs["computed_locally"] = "true"
                        except Exception as exc:  # 无法求值时保留公式并提示
                            res.warnings.append(f"{ws_f.title}!{coord} 公式无法本地求值：{exc}")
                text = _fmt(val)
                out_row.append(text)
                if c.row in hidden_rows or c.column in hidden_cols:
                    if text:
                        res.hidden.append(HiddenContent("隐藏行列", f"{ws_f.title}!{coord}={text}"))
                    continue
                if text:
                    u = b.add("sheet_cell", text, f"{ws_f.title}!{coord}", **attrs)
                    if c.comment is not None and c.comment.text.strip():
                        cu = b.add("comment", c.comment.text.strip(), f"{ws_f.title}!{coord}#comment", sheet=ws_f.title)
                        res.relations.append(SourceRelation(kind="note_of", src=cu.unit_id, dst=u.unit_id))
                        res.hidden.append(HiddenContent("单元格批注", c.comment.text.strip()))
            grid.append(out_row)
        table = table_from_rows(material_id, tid, grid, ws_f.title, title=ws_f.title)
        # 表注（如“注：仅统计已验收项目”）
        for r in list(table.rows):
            joined = "".join(r).strip()
            if joined and NOTE_RE.match(joined):
                table.notes.append(joined)
        res.tables.append(table)
    return res
